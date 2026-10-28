"""One complete table of the library: CSV, Markdown or XLSX."""
from __future__ import annotations
import csv
import datetime as dt
import io
from typing import Any

from .store import Store

KIND_LABELS = {"work": "作品", "prompt": "提示词", "material": "素材",
               "record": "记录", "skill": "Skills", "note": "便签"}
COLUMNS = ("类型", "分组", "编号", "标题", "客户", "日期", "版本", "提示词/正文",
           "说明", "标签", "工具", "状态", "发布平台", "发布链接", "文件", "路径",
           "更新时间")
# 数据记录: one row per observation (kind "metrics" / the 「数据」 sheet)
METRICS_KIND = "metrics"
METRICS_LABEL = "数据"
METRICS_COLUMNS = ("作品", "客户", "平台", "账号", "记录时间", "距发布小时", "节点", "浏览",
                   "触达", "点赞", "评论", "收藏", "分享", "主页访问", "涨粉", "点击率",
                   "三秒率", "完播率", "平均观看秒", "来源", "备注")
METRIC_FIELDS = ("views", "reach", "likes", "comments", "saves", "shares",
                 "profile_visits", "follows", "ctr", "three_sec_rate",
                 "completion_rate", "avg_watch_sec")
FORMATS = {"csv": ("text/csv; charset=utf-8", "csv"),
           "md": ("text/markdown; charset=utf-8", "md"),
           "xlsx": ("application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet", "xlsx")}
STATUS_LABELS = {"published": "已发布", "ready": "待发布", "cut": "已剪辑",
                 "wip": "制作中"}
WIDE = {"提示词/正文": 60, "说明": 30, "标题": 28, "路径": 40, "发布链接": 36}


def stamp(value: str | None) -> str:
    try:
        moment = dt.datetime.fromisoformat(value or "")
    except ValueError:
        return value or ""
    return moment.astimezone().strftime("%Y-%m-%d %H:%M")


class ExportError(ValueError):
    """The requested export cannot be produced."""


def work_row(work: dict[str, Any]) -> dict[str, str]:
    return {"类型": "作品", "分组": work["category"], "标题": work["title"],
            "客户": work["client"], "日期": work["date"],
            "版本": work["version"], "提示词/正文": work["prompt"],
            "说明": work.get("note") or "", "标签": work.get("tags") or "",
            "工具": work.get("tool") or "",
            "状态": STATUS_LABELS.get(work["status"], work["status"]),
            "发布平台": work["platforms"], "发布链接": work["urls"],
            "文件": work.get("cover_source") or "",
            "路径": work["relative_path"], "更新时间": stamp(work["updated_at"])}


def item_row(item: dict[str, Any]) -> dict[str, str]:
    return {"类型": KIND_LABELS.get(item["kind"], item["kind"]),
            "分组": item["section"], "编号": item.get("code") or "",
            "标题": item["title"], "客户": item.get("client") or "",
            "版本": item.get("version") or "",
            "提示词/正文": item["text"],
            "说明": item["note"], "标签": item["tags"], "工具": item["tool"],
            "状态": item.get("status") or "", "文件": item.get("file") or "",
            "路径": item["path"], "更新时间": stamp(item["modified_at"])}


def rows(store: Store, kind: str = "",
         all_versions: bool = False) -> list[dict[str, str]]:
    """The rows the list shows: one version of each work's prompt unless all_versions."""
    if kind and kind not in KIND_LABELS:
        raise ExportError("未知的导出类型")
    result = []
    if kind in ("", "work"):
        result += [work_row(work) for work in store.export_works()]
    if kind != "work":
        result += [item_row(item) for item in
                   store.list_items(kind, all_versions=all_versions)]
    return [{column: str(row.get(column) or "") for column in COLUMNS}
            for row in result]


def metric_text(value: int | float | None) -> str:
    return "" if value is None else str(value)


def metrics_rows(store: Store) -> list[dict[str, str]]:
    """One row per observation of every post, newest post first; 节点 is the
    D1/D3/D7 the observation stands for, if any."""
    result = []
    for post in store.metrics_posts():
        for point in post["history"]:
            row = {"作品": post["title"], "客户": post["client"], "平台": post["platform"],
                   "账号": post["account"], "记录时间": stamp(point["observed_at"]),
                   "距发布小时": metric_text(point["age_hours"]),
                   "节点": point["window"] or "", "来源": point["source"],
                   "备注": point["note"]}
            for column, field in zip(METRICS_COLUMNS[7:19], METRIC_FIELDS):
                row[column] = metric_text(point["metrics"].get(field))
            result.append({column: row.get(column, "") for column in METRICS_COLUMNS})
    return result


def sheets(data: list[dict[str, str]], store: Store | None,
           kind: str) -> list[tuple[str, tuple[str, ...], list[dict[str, str]]]]:
    """The tables an export holds: the asset table (per kind for xlsx/md), then
    the 数据 table when everything or only metrics is asked for."""
    if kind == METRICS_KIND:
        return [(METRICS_LABEL, METRICS_COLUMNS, data)]
    result = [("全部", COLUMNS, data)] + [
        (label, COLUMNS, [row for row in data if row["类型"] == label])
        for label in KIND_LABELS.values()]
    if not kind and store is not None:
        result.append((METRICS_LABEL, METRICS_COLUMNS, metrics_rows(store)))
    return result


def to_csv(data: list[dict[str, str]], columns: tuple[str, ...] = COLUMNS) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, columns)
    writer.writeheader()
    writer.writerows(data)
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def md_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def to_markdown(tables: list[tuple[str, tuple[str, ...], list[dict[str, str]]]]) -> bytes:
    lines = ["# SPT 资产总表", ""]
    for label, all_columns, group in tables:
        if not group or label == "全部":
            continue
        columns = [column for column in all_columns
                   if column != "类型" and any(row[column] for row in group)]
        lines += [f"## {label}（{len(group)}）", "",
                  "| " + " | ".join(columns) + " |",
                  "|" + "---|" * len(columns)]
        lines += ["| " + " | ".join(md_cell(row[column]) for column in columns)
                  + " |" for row in group]
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def to_xlsx(tables: list[tuple[str, tuple[str, ...], list[dict[str, str]]]]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise ExportError("当前 Python 未安装 openpyxl，请改用 CSV 导出") from error
    book = Workbook()
    book.remove(book.active)
    for title, columns, group in tables:
        if not group and title not in ("全部", METRICS_LABEL):
            continue
        sheet = book.create_sheet(title)
        sheet.append(columns)
        for row in group:
            sheet.append([clean(row[column]) for column in columns])
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(1, index)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EEF1F6")
            sheet.column_dimensions[get_column_letter(index)].width = \
                WIDE.get(column, 14)
        wrap = Alignment(wrap_text=True, vertical="top")
        for line in sheet.iter_rows(min_row=2):
            for cell in line:
                cell.alignment = wrap
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def clean(value: str) -> str:
    text = "".join(char for char in value
                   if char in "\t\n" or ord(char) >= 32)[:32000]
    return " " + text if text.startswith("=") else text


def export(store: Store, format_name: str, kind: str = "",
           day: dt.date | None = None,
           all_versions: bool = False) -> tuple[bytes, str, str]:
    """(content, mime type, file name); the name reads like SPT-资产-2026-10-22.xlsx,
    or SPT-提示词-2026-10-22.xlsx when only one kind is exported."""
    if format_name not in FORMATS:
        raise ExportError("导出格式只支持 csv、xlsx、md")
    if kind == METRICS_KIND:
        data, columns = metrics_rows(store), METRICS_COLUMNS
    else:
        data, columns = rows(store, kind, all_versions), COLUMNS
    if format_name == "csv":
        content = to_csv(data, columns)
    else:
        tables = sheets(data, store, kind)
        content = to_markdown(tables) if format_name == "md" else to_xlsx(tables)
    mime, suffix = FORMATS[format_name]
    label = METRICS_LABEL if kind == METRICS_KIND else KIND_LABELS.get(kind, "资产")
    return content, mime, f"SPT-{label}-{(day or dt.date.today()).isoformat()}.{suffix}"
