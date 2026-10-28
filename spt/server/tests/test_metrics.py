"""数据记录: files in, index, queries, writes that only append, dashboard merge."""
from __future__ import annotations
import contextlib
import csv
import datetime as dt
import http.client
import io
import json
import tempfile
import threading
import unittest
from pathlib import Path
from typing import Any
from unittest.mock import patch

from .. import agent, cli
from ..api import Server
from ..exporter import METRICS_COLUMNS, export
from ..indexer import Indexer
from ..mcp_server import TOOLS, call
from ..metrics import (assign_nodes, parse_time, platform_code, read_analysis,
                       read_performance, split_header)
from ..metrics_dashboard import merge_workspace_metrics
from ..metrics_write import MetricsError, WorkLookup, record_observation
from .support import fixture_config, populate, write

CAR = "运营/2026-09-08-示例创作者-小车先停好"
RECORD = ("# 当前采用：V1（2026-09-08）\n"
          "- 小红书: https://www.xiaohongshu.com/explore/abc123；2026-09-08 17:22 PDT，已核验\n")
PERFORMANCE = (
    "记录时间,平台,账号,浏览,触达,点赞,评论,收藏,分享,主页访问,涨粉,点击率,三秒率,完播率,平均观看秒,来源,备注\n"
    "2026-09-09T20:00:00-07:00,rednote,示例创作者,60,,6,1,1,0,,,,,,,后台,\n"
    "2026-09-10T03:35:00-07:00,小红书,示例创作者,72,,8,1,1,1,,,,,,,原始快照/a.json,\n"
    "2026-09-11T21:35:00-07:00,rednote,示例创作者,82,,8,2,1,1,,,,,41.5%,,原始快照/b.json,D3 晚到 4.2 小时\n"
    "2026-09-16T00:00:00-07:00,rednote,示例创作者,90,,9,2,1,1,,,,,,,,\n"
    "2026-09-20,rednote,示例创作者,92,,9,2,2,1,,,,,,,,补录\n"
    "昨天,rednote,示例创作者,1,,,,,,,,,,,,,\n"
    "2026-09-21T00:00:00-07:00,rednote,示例创作者,abc,,,,,,,,,,,,,\n"
    "2026-09-22T00:00:00-07:00,火星,示例创作者,5,,,,,,,,,,,,,\n")
RETENTION = ("记录时间,平台,账号,秒,留存率\n"
             "2026-09-10T03:35:00-07:00,rednote,示例创作者,0,100\n"
             "2026-09-10T03:35:00-07:00,rednote,示例创作者,3,71.5\n"
             "2026-09-11T21:35:00-07:00,rednote,示例创作者,0,100\n"
             "2026-09-11T21:35:00-07:00,rednote,示例创作者,3,58\n"
             "2026-09-11T21:35:00-07:00,rednote,示例创作者,10,40\n")
ANALYSIS_D3 = ("平台: rednote\n账号: 示例创作者\n节点: D3\n结论: 开头 3 秒没有冲突，封面和标题重复\n"
               "标签: 开头弱, 封面\n\n## 现象\n前 3 秒留存从 100 掉到 58\n")
ANALYSIS_WEEK = "---\n节点: 周\n结论: 发布时间太晚\n标签: 发布时间、开头弱\n---\n正文在这里\n"
ACCOUNT = ("记录时间,粉丝,浏览,触达,互动,点赞,评论,收藏,分享,主页访问,涨粉,来源,备注\n"
           "2026-10-21T03:48:26-07:00,200,,,,,,,,,,公开主页,\n"
           "2026-10-14T03:48:26-07:00,190,,,,,,,,,,公开主页,\n")
CLIENTS = ("# 客户\n\n| 简称 | 名字 | 类型 | 颜色 |\n|---|---|---|---|\n"
           "| 示例创作者 | 示例创作者 | 自己 | |\n| AK | Sample Studio | 客户 | |\n")


def build(workspace: Path) -> dict[str, Path]:
    """The shared fixture works plus one with a full 数据/ folder and account-level data."""
    paths = populate(workspace)
    car = workspace / CAR
    write(car / "发布记录.md", RECORD)
    write(car / "数据/表现.csv", PERFORMANCE)
    write(car / "数据/留存.csv", RETENTION)
    write(car / "数据/分析/2026-09-12-D3.md", ANALYSIS_D3)
    write(car / "数据/分析/周复盘.md", ANALYSIS_WEEK)
    # a work without a publish time: a timeline, no nodes
    write(paths["火光"] / "数据/表现.csv",
          "记录时间,平台,账号,浏览,点赞\n2026-10-21T10:00:00Z,instagram,,30,2\n"
          "2026-10-22T10:00:00Z,tiktok,ak_tiktok,40,\n")
    write(workspace / "数据/账号/rednote-示例创作者/表现.csv", ACCOUNT)
    write(workspace / "数据/账号/rednote-示例创作者/分析/2026-10-20-周复盘.md",
          "结论: 周更不够\n标签: 发布时间\n\n本周只发了一篇。\n")
    write(workspace / "数据/账号/not-a-platform/表现.csv", ACCOUNT)
    write(workspace / "客户.md", CLIENTS)
    paths["小车"] = car
    return paths


def snapshot_accounts() -> dict[str, Any]:
    """A dashboard snapshot like dashboard.py builds, before the merge."""
    def account(identifier: str, brand: str, platform: str, handle: str | None) -> dict[str, Any]:
        return {"id": identifier, "brand_id": identifier.split(":")[0], "brand": brand,
                "name": brand, "platform": platform, "platform_label": platform,
                "handle": handle, "data_state": "unavailable", "as_of": None,
                "metrics": {"followers": None, "views": None}, "history": [],
                "posts_count": 0, "login_state": "not_checked"}
    return {"accounts": [
        account("parenting:rednote", "示例创作者", "rednote", "https://www.rednote.com/user/x"),
        account("demo-brand:instagram", "Sample Studio", "instagram", "test_kitchen"),
        account("demo-brand:facebook", "Sample Studio", "facebook", None)],
        "posts": [], "capabilities": {}}


class ParsingTests(unittest.TestCase):
    def test_time_forms(self) -> None:
        utc = dt.timezone.utc
        self.assertEqual(parse_time("2026-09-08 17:22 PDT").astimezone(utc).isoformat(),
                         "2026-09-09T00:22:00+00:00")
        self.assertEqual(parse_time("2026-01-08 17:22 PST").utcoffset(), dt.timedelta(hours=-8))
        self.assertEqual(parse_time("2026-09-08T10:00:00Z").utcoffset(), dt.timedelta(0))
        self.assertEqual(parse_time("2026-09-08 10:00 UTC").utcoffset(), dt.timedelta(0))
        self.assertEqual(parse_time("2026-10-22T09:30:00+08:00").utcoffset(), dt.timedelta(hours=8))
        self.assertEqual(parse_time("2026-10-22T09:30:00-0700").utcoffset(), dt.timedelta(hours=-7))
        naive = parse_time("2026-10-22 09:30")
        self.assertEqual(naive, dt.datetime(2026, 10, 22, 9, 30).astimezone())
        self.assertEqual(parse_time("2026-10-22"), dt.datetime(2026, 10, 22, 12, 0).astimezone())
        for bad in ("", "昨天", "2026-13-01", "10/22/2026", "2026-10-22 09:30 XYZ"):
            with self.subTest(bad=bad):
                self.assertIsNone(parse_time(bad))

    def test_platform_vocabulary(self) -> None:
        for text, code in (("rednote", "rednote"), ("小红书", "rednote"), ("xiaohongshu", "rednote"),
                           ("抖音", "douyin"), ("视频号", "wechat_channels"), ("B站", "bilibili"),
                           ("油管", "youtube"), ("Instagram", "instagram"),
                           ("google_business_profile", "google_business"),
                           ("Google 商家资料", "google_business"), ("Yelp", "yelp"),
                           ("火星", None), ("", None)):
            with self.subTest(text=text):
                self.assertEqual(platform_code(text), code)

    def test_csv_rows_missing_cells_percent_and_bad_rows(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "表现.csv"
            write(path, PERFORMANCE)
            issues: list[str] = []
            rows = read_performance(path, "x/数据/表现.csv", issues)
        self.assertEqual([row.views for row in rows], [60, 72, 82, 90, 92])
        self.assertEqual(rows[0].at, "2026-09-10T03:00:00+00:00")
        self.assertEqual(rows[1].platform, "rednote")          # 小红书
        self.assertIsNone(rows[0].reach)                       # empty is unknown
        self.assertEqual(rows[0].shares, 0)                    # 0 is 0
        self.assertEqual(rows[2].completion_rate, 41.5)        # 41.5%
        self.assertEqual(rows[2].note, "D3 晚到 4.2 小时")
        self.assertEqual(issues, [
            "x/数据/表现.csv 第 7 行：记录时间「昨天」看不懂，已跳过",
            "x/数据/表现.csv 第 8 行：浏览「abc」不是数字，已跳过",
            "x/数据/表现.csv 第 9 行：平台「火星」不认识，已跳过"])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "表现.csv"
            write(path, "平台,浏览,多余的列\nrednote,5,x\n")
            issues = []
            self.assertEqual(read_performance(path, "y", issues), [])
            self.assertEqual(issues, ["y：表头里没有「记录时间」"])
            write(path, "浏览,记录时间,平台,多余的列\n5,2026-10-01T00:00:00Z,rednote,x\n")
            [row] = read_performance(path, "y", [])
            self.assertEqual((row.views, row.account, row.likes), (5, "", None))

    def test_node_assignment_prefers_closest_and_reports_lateness(self) -> None:
        def point(age: float | None, at: str) -> dict[str, Any]:
            return {"age_hours": age, "at": at, "metrics": {"views": age},
                    "platform": "rednote", "account": ""}
        nodes = assign_nodes([point(17.9, "a"), point(26.63, "b"), point(34.2, "c"),
                              point(76.22, "d"), point(119.9, "e"), point(174.63, "f"),
                              point(240, "g")])
        self.assertEqual((nodes["D1"]["at"], nodes["D1"]["late_hours"]), ("b", 2.63))
        self.assertEqual((nodes["D3"]["at"], nodes["D3"]["late_hours"]), ("d", 4.22))
        self.assertEqual((nodes["D7"]["at"], nodes["D7"]["late_hours"]), ("f", 6.63))
        early = assign_nodes([point(20, "a"), point(23, "b"), point(23, "c")])
        self.assertEqual((early["D1"]["at"], early["D1"]["late_hours"]), ("b", -1))
        self.assertIsNone(early["D3"])
        self.assertEqual(assign_nodes([point(None, "a")]), {"D1": None, "D3": None, "D7": None})

    def test_node_prefers_a_reading_with_views(self) -> None:
        public = {"age_hours": 172.2, "at": "a", "metrics": {"views": None, "likes": 8},
                  "platform": "rednote", "account": ""}
        backend = {"age_hours": 173.6, "at": "b", "metrics": {"views": 92, "likes": 8},
                   "platform": "rednote", "account": ""}
        self.assertEqual(assign_nodes([public, backend])["D7"]["at"], "b")
        self.assertEqual(assign_nodes([public])["D7"]["at"], "a")

    def test_analysis_header_fenced_bare_tags_and_date_fallback(self) -> None:
        header, body = split_header(ANALYSIS_D3)
        self.assertEqual(header["节点"], "D3")
        self.assertEqual(body.splitlines()[0], "## 现象")
        header, body = split_header(ANALYSIS_WEEK)
        self.assertEqual((header["结论"], body), ("发布时间太晚", "正文在这里"))
        self.assertEqual(split_header("## 只有正文\n开头: 不是键值\n"), ({}, "## 只有正文\n开头: 不是键值"))
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            write(root / "2026-09-12-D3.md", ANALYSIS_D3)
            write(root / "周复盘.md", ANALYSIS_WEEK)
            write(root / "2026-01-01-x.md", "日期: 2026-02-02\n结论: 头里的日期优先\n")
            first = read_analysis(root / "2026-09-12-D3.md", "a.md")
            self.assertEqual((first.date, first.platform, first.account, first.node),
                             ("2026-09-12", "rednote", "示例创作者", "D3"))
            self.assertEqual(first.tags, ["开头弱", "封面"])
            week = read_analysis(root / "周复盘.md", "b.md", "rednote", "示例创作者")
            self.assertEqual(week.tags, ["发布时间", "开头弱"])
            self.assertEqual(week.date, dt.date.today().isoformat())
            self.assertEqual((week.platform, week.account), ("rednote", "示例创作者"))
            self.assertEqual(read_analysis(root / "2026-01-01-x.md", "c.md").date, "2026-02-02")


class IndexTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.config = fixture_config(self.root)
        self.paths = build(self.config.workspace)
        self.indexer = Indexer(self.config)
        self.result = self.indexer.reindex()
        self.car = next(work for work in self.indexer.store.list_works()
                        if work["title"] == "小车先停好")

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_scan_reads_the_folder_and_reports_bad_rows_in_health(self) -> None:
        self.assertEqual(self.result["count"], 5)           # 数据/ at the root is not a 专业
        health = self.indexer.health()
        self.assertTrue(health["capabilities"]["metrics"])
        self.assertIn(f"{CAR}/数据/表现.csv 第 7 行：记录时间「昨天」看不懂，已跳过", health["issues"])
        self.assertIn(f"{CAR}/数据/表现.csv 第 8 行：浏览「abc」不是数字，已跳过", health["issues"])
        self.assertIn("数据/账号/not-a-platform：文件夹名要是「平台-账号」，已跳过", health["issues"])
        self.assertNotIn("数据", {work["category"] for work in self.indexer.store.list_works()})

    def test_work_metrics_nodes_latest_retention_and_analyses(self) -> None:
        metrics = self.indexer.store.work_metrics(self.car["id"])
        self.assertEqual([item["metrics"]["views"] for item in metrics["observations"]],
                         [60, 72, 82, 90, 92])
        self.assertEqual(metrics["observations"][0]["age_hours"], 26.63)
        nodes = metrics["nodes"]
        self.assertEqual((nodes["D1"]["metrics"]["views"], nodes["D1"]["late_hours"]), (60, 2.63))
        self.assertEqual((nodes["D3"]["metrics"]["views"], nodes["D3"]["late_hours"]), (82, 4.22))
        self.assertEqual((nodes["D7"]["metrics"]["views"], nodes["D7"]["late_hours"]), (90, 6.63))
        self.assertEqual(nodes["D3"]["metrics"]["completion_rate"], 41.5)
        self.assertEqual(metrics["latest"]["metrics"]["views"], 92)
        self.assertEqual(metrics["retention"][0]["points"], [[0, 100], [3, 58], [10, 40]])
        self.assertEqual(len(metrics["retention"]), 2)
        self.assertEqual([item["node"] for item in metrics["analyses"]], ["周", "D3"])
        self.assertEqual(metrics["analyses"][1]["tags"], ["开头弱", "封面"])
        self.assertIn("## 现象", metrics["analyses"][1]["body"])
        self.assertEqual(metrics["analyses"][1]["path"], f"{CAR}/数据/分析/2026-09-12-D3.md")
        [post] = metrics["posts"]
        self.assertEqual((post["platform"], post["account"], post["metric_window"]),
                         ("rednote", "示例创作者", None))
        self.assertEqual(post["published_at"], "2026-09-08T17:22:00-07:00")
        self.assertEqual(post["url"], "https://www.xiaohongshu.com/explore/abc123")
        self.assertEqual(self.indexer.store.detail(self.car["id"])["metrics"]["latest"],
                         metrics["latest"])

    def test_no_publish_time_means_timeline_only(self) -> None:
        fire = next(work for work in self.indexer.store.list_works() if work["title"] == "火光")
        metrics = self.indexer.store.work_metrics(fire["id"])
        self.assertEqual([item["age_hours"] for item in metrics["observations"]], [None, None])
        self.assertEqual(metrics["nodes"], {"D1": None, "D3": None, "D7": None})
        self.assertEqual(len(metrics["posts"]), 2)        # instagram and tiktok
        self.assertIsNone(metrics["posts"][0]["published_at"])

    def test_posts_filters_tags_and_accounts(self) -> None:
        store = self.indexer.store
        self.assertEqual(len(store.metrics_posts()), 3)
        posts = store.metrics_posts(client="示例创作者")
        self.assertEqual([post["title"] for post in posts], ["小车先停好"])
        self.assertEqual(posts[0]["source_kind"], "workspace_metrics")
        self.assertEqual(posts[0]["history"][0]["window"], "D1")
        self.assertEqual(posts[0]["history"][2]["window"], "D3")
        self.assertIsNone(posts[0]["history"][1]["window"])
        self.assertEqual(posts[0]["metrics_by_window"]["D7"]["views"], 90)
        self.assertIsNone(posts[0]["metrics_by_window"]["D1"]["reach"])
        self.assertEqual(posts[0]["metric_observed_at_by_window"]["D1"], "2026-09-10T03:00:00+00:00")
        self.assertEqual(posts[0]["retention"]["at"], "2026-09-12T04:35:00+00:00")
        self.assertNotIn("body", posts[0]["analyses"][0])
        self.assertEqual(posts[0]["tags"], ["发布时间", "封面", "开头弱"])
        self.assertEqual(len(store.metrics_posts(platform="小红书")), 1)
        self.assertEqual(len(store.metrics_posts(platform="tiktok", account="ak_tiktok")), 1)
        self.assertEqual(len(store.metrics_posts(tag="封面")), 1)
        self.assertEqual(len(store.metrics_posts(tag="没有的")), 0)
        self.assertEqual(len(store.metrics_posts(since="2026-09-01", until="2026-09-30")), 1)
        self.assertEqual(len(store.metrics_posts(since="2026-10-22")), 1)   # 火光 tiktok by latest
        self.assertEqual(len(store.metrics_posts(q="封面 小车")), 1)
        with self.assertRaises(ValueError):
            store.metrics_posts(platform="火星")
        tags = store.metrics_tags()
        self.assertEqual(tags[0], {"tag": "发布时间", "count": 2, "post_ids": [posts[0]["id"]]})
        self.assertEqual({item["tag"]: item["count"] for item in tags},
                         {"发布时间": 2, "开头弱": 2, "封面": 1})
        self.assertEqual([item["tag"] for item in store.metrics_tags(client="示例创作者")],
                         ["开头弱", "发布时间", "封面"])
        [account] = store.account_metrics()
        self.assertEqual((account["platform"], account["account"]), ("rednote", "示例创作者"))
        self.assertEqual(account["latest"]["followers"], 200)
        self.assertEqual([point["metrics"]["followers"] for point in account["history"]], [190, 200])
        self.assertEqual(account["analyses"][0]["conclusion"], "周更不够")

    def test_dashboard_merge_attaches_or_adds_accounts(self) -> None:
        snapshot = snapshot_accounts()
        clients = self.indexer.clients.read(self.indexer.workspace)[0]
        merge_workspace_metrics(snapshot, self.indexer.store, clients)
        accounts = {item["id"]: item for item in snapshot["accounts"]}
        parenting = accounts["parenting:rednote"]
        car = next(post for post in snapshot["posts"] if post["title"] == "小车先停好")
        self.assertEqual(car["account_id"], "parenting:rednote")
        self.assertEqual(parenting["posts_count"], 1)
        self.assertEqual(parenting["data_state"], "workspace")
        # account-level rows filled followers and history
        self.assertEqual(parenting["metrics"]["followers"], 200)
        self.assertEqual(parenting["as_of"], "2026-10-21T10:48:26+00:00")
        self.assertEqual(len(parenting["history"]), 2)
        fire = next(post for post in snapshot["posts"]
                    if post["title"] == "火光" and post["platform"] == "instagram")
        self.assertEqual(fire["account_id"], "demo-brand:instagram")
        # the one post with data rows plus the publish records that have none yet (2026-10-24)
        self.assertEqual(accounts["demo-brand:instagram"]["posts_count"], 3)
        new = accounts["AK:tiktok:ak_tiktok"]
        self.assertEqual((new["brand"], new["brand_id"], new["handle"], new["data_state"],
                          new["source"], new["posts_count"], new["login_state"]),
                         ("Sample Studio", "AK", "ak_tiktok", "workspace", "工作区/数据", 1,
                          "not_checked"))
        self.assertEqual(snapshot["analysis_tags"][0]["tag"], "发布时间")
        self.assertEqual(car["analyses"][1]["conclusion"], "开头 3 秒没有冲突，封面和标题重复")
        self.assertEqual(car["retention"]["points"][1], [3, 58])
        self.assertTrue(snapshot["capabilities"]["workspace_metrics"])
        self.assertEqual(snapshot["posts"][0]["title"], "火光")   # newest first

    def test_export_has_a_metrics_table(self) -> None:
        data, _, name = export(self.indexer.store, "csv", "metrics", dt.date(2026, 10, 22))
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(tuple(rows[0]), METRICS_COLUMNS)
        self.assertEqual(name, "SPT-数据-2026-10-22.csv")
        self.assertEqual(len(rows), 1 + 5 + 2)
        car = [row for row in rows[1:] if row[0] == "小车先停好"]
        self.assertEqual([row[6] for row in car], ["D1", "", "D3", "D7", ""])
        self.assertEqual((car[2][17], car[2][20]), ("41.5", "D3 晚到 4.2 小时"))
        try:
            import openpyxl
        except ImportError:
            return
        book = openpyxl.load_workbook(io.BytesIO(export(self.indexer.store, "xlsx")[0]))
        self.assertIn("数据", book.sheetnames)
        self.assertEqual(book["数据"].cell(1, 7).value, "节点")

    def test_cli_add_writes_the_file_then_show_sees_it(self) -> None:
        table = self.paths["小车"] / "数据/表现.csv"
        before = table.read_text(encoding="utf-8")
        with patch.object(agent, "open_indexer", return_value=self.indexer):
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                code = cli.main(["metrics", "add", "--work", "小车先停好", "--platform", "小红书",
                                 "--at", "2026-09-25T10:00:00-07:00", "--views", "100",
                                 "--completion-rate", "45%", "--source", "后台", "--json"])
            self.assertEqual(code, 0)
            written = json.loads(out.getvalue())
            self.assertEqual(written["path"], f"{CAR}/数据/表现.csv")
            self.assertEqual(written["row"]["metrics"]["views"], 100)
            self.assertEqual(written["index"]["refreshed"], self.car["id"])
            after = table.read_text(encoding="utf-8")
            self.assertTrue(after.startswith(before))
            self.assertEqual(after[len(before):],
                             "2026-09-25T10:00:00-07:00,rednote,,100,,,,,,,,,,45,,后台,\n")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                self.assertEqual(cli.main(["metrics", "show", "小车先停好", "--json"]), 0)
            shown = json.loads(out.getvalue())
            self.assertEqual(shown["latest"]["metrics"]["views"], 100)
            self.assertEqual(len(shown["observations"]), 6)
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                self.assertEqual(cli.main(["metrics", "--client", "示例创作者"]), 0)
            self.assertIn("小车先停好", out.getvalue())
            self.assertIn("D1", out.getvalue().splitlines()[0])
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                self.assertEqual(cli.main(["metrics", "tags", "--client", "示例创作者", "--json"]), 0)
            self.assertEqual(json.loads(out.getvalue())[0]["tag"], "开头弱")
            err = io.StringIO()
            with contextlib.redirect_stderr(err):
                self.assertEqual(cli.main(["metrics", "add", "--work", "没有这个", "--platform",
                                           "rednote", "--views", "1"]), 1)
            self.assertIn("找不到作品", err.getvalue())

    def test_writes_validate_and_never_overwrite(self) -> None:
        store = self.indexer.store
        with self.assertRaises(WorkLookup):
            record_observation(self.indexer, {"work": "早期", "platform": "rednote", "views": 1})
        for value in ({"work": "小车", "platform": "火星", "views": 1},
                      {"work": "小车", "platform": "rednote"},
                      {"work": "小车", "platform": "rednote", "views": "abc"},
                      {"work": "小车", "platform": "rednote", "views": -1},
                      {"work": "小车", "platform": "rednote", "views": 1, "at": "昨天"},
                      {"work": "小车", "platform": "rednote", "views": 1, "account": "a/b"}):
            with self.subTest(value=value), self.assertRaises(MetricsError):
                record_observation(self.indexer, value)
        with patch.object(agent, "open_indexer", return_value=self.indexer):
            first = call("spt_metrics_analyse", {"work": "小车", "platform": "rednote", "node": "D7",
                                                 "conclusion": "结尾没有引导", "tags": ["结尾无引导"],
                                                 "body": "## 现象\n看完就走", "date": "2026-09-16"})
            second = call("spt_metrics_analyse", {"work": "小车", "node": "D7",
                                                  "conclusion": "再看一次", "date": "2026-09-16"})
        self.assertEqual(first["path"], f"{CAR}/数据/分析/2026-09-16-D7.md")
        self.assertEqual(second["path"], f"{CAR}/数据/分析/2026-09-16-D7-2.md")
        text = (self.config.workspace / first["path"]).read_text(encoding="utf-8")
        self.assertEqual(text, "平台: rednote\n节点: D7\n结论: 结尾没有引导\n标签: 结尾无引导\n\n"
                               "## 现象\n看完就走\n")
        analyses = store.work_metrics(self.car["id"])["analyses"]
        self.assertEqual([item["node"] for item in analyses], ["周", "D7", "D7", "D3"])
        self.assertEqual(len(store.metrics_tags(tag="结尾无引导")), 1)
        # a work folder swapped for a link is refused
        folder = self.paths["小车"]
        folder.rename(folder.with_name("backup"))
        folder.symlink_to(folder.with_name("backup"), target_is_directory=True)
        with self.assertRaises(ValueError):
            record_observation(self.indexer, {"work": "小车", "platform": "rednote", "views": 1})

    def test_mcp_tools_read_and_write(self) -> None:
        names = [tool["name"] for tool in TOOLS]
        self.assertEqual(names[-3:], ["spt_metrics", "spt_metrics_add", "spt_metrics_analyse"])
        with patch.object(agent, "open_indexer", return_value=self.indexer):
            posts = call("spt_metrics", {"client": "示例创作者", "platform": "小红书"})
            self.assertEqual([post["title"] for post in posts], ["小车先停好"])
            shown = call("spt_metrics", {"work": "2026-09-08-示例创作者-小车先停好"})
            self.assertEqual(shown["nodes"]["D3"]["late_hours"], 4.22)
            self.assertEqual(call("spt_metrics", {"tags": True})[0]["tag"], "发布时间")
            added = call("spt_metrics_add", {"kind": "account", "platform": "小红书",
                                             "account": "示例创作者", "followers": 210,
                                             "at": "2026-10-28T10:00:00-07:00", "source": "主页"})
            self.assertEqual(added["path"], "数据/账号/rednote-示例创作者/表现.csv")
            self.assertEqual(added["index"]["refreshed"], "accounts")
            [account] = self.indexer.store.account_metrics()
            self.assertEqual(account["latest"]["followers"], 210)
            curve = call("spt_metrics_add", {"kind": "retention", "work": "小车", "platform": "rednote",
                                             "account": "示例创作者", "points": [[0, 100], [3, 66]],
                                             "at": "2026-09-30T10:00:00-07:00"})
            self.assertEqual(curve["points"], [[0, 100], [3, 66]])
            self.assertEqual(self.indexer.store.work_metrics(self.car["id"])["retention"][0]["points"],
                             [[0, 100], [3, 66]])
            lines = (self.paths["小车"] / "数据/留存.csv").read_text(encoding="utf-8").splitlines()
            self.assertEqual(lines[-1], "2026-09-30T10:00:00-07:00,rednote,示例创作者,3,66")
            new_account = call("spt_metrics_add", {"kind": "account", "platform": "douyin",
                                                   "account": "新号", "followers": 1})
            self.assertEqual(new_account["path"], "数据/账号/douyin-新号/表现.csv")
            text = (self.config.workspace / new_account["path"]).read_text(encoding="utf-8")
            self.assertTrue(text.startswith("记录时间,粉丝,浏览,触达,互动,点赞,评论,收藏,分享,主页访问,涨粉,来源,备注\n"))


class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temporary = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temporary.name)
        cls.config = fixture_config(cls.root)
        cls.paths = build(cls.config.workspace)
        cls.indexer = Indexer(cls.config)
        cls.indexer.reindex()
        cls.server = Server(cls.indexer)
        cls.port = cls.server.server_port
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.car = next(work for work in cls.indexer.store.list_works()
                       if work["title"] == "小车先停好")

    @classmethod
    def tearDownClass(cls) -> None:
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=3)
        cls.temporary.cleanup()

    def get(self, path: str) -> tuple[int, Any]:
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=3)
        connection.request("GET", path)
        response = connection.getresponse()
        result = response.status, json.loads(response.read() or b"null")
        connection.close()
        return result

    def post(self, path: str, value: object, content_type: str = "application/json") -> tuple[int, Any]:
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=3)
        connection.request("POST", path, body=json.dumps(value).encode(),
                           headers={"Content-Type": content_type})
        response = connection.getresponse()
        result = response.status, json.loads(response.read() or b"{}")
        connection.close()
        return result

    def test_reads(self) -> None:
        status, body = self.get("/api/metrics?client=%E7%A4%BA%E4%BE%8B%E5%88%9B%E4%BD%9C%E8%80%85&platform=rednote")
        self.assertEqual(status, 200)
        self.assertEqual(set(body), {"posts", "tags", "accounts"})
        self.assertEqual(body["posts"][0]["title"], "小车先停好")
        self.assertEqual(body["accounts"][0]["latest"]["followers"], 200)
        status, body = self.get(f"/api/works/{self.car['id']}/metrics")
        self.assertEqual(status, 200)
        self.assertEqual(body["nodes"]["D1"]["late_hours"], 2.63)
        self.assertEqual(self.get("/api/works/" + "0" * 20 + "/metrics")[0], 404)
        self.assertEqual(self.get("/api/metrics?platform=%E7%81%AB%E6%98%9F")[0], 400)
        status, detail = self.get(f"/api/works/{self.car['id']}")
        self.assertEqual(detail["metrics"]["latest"]["metrics"]["views"], 92)
        status, dashboard = self.get("/api/dashboard")
        self.assertEqual(status, 200)
        self.assertIn("analysis_tags", dashboard)
        self.assertTrue(any(post["source_kind"] == "workspace_metrics" for post in dashboard["posts"]))

    def test_writes_round_trip(self) -> None:
        status, body = self.post("/api/metrics/add", {"work": "小车先停好", "platform": "rednote",
                                                      "account": "示例创作者", "views": 120, "likes": 10,
                                                      "at": "2026-09-26T10:00:00-07:00",
                                                      "note": "接口"})
        self.assertEqual(status, 200, body)
        self.assertEqual((body["work_id"], body["path"]), (self.car["id"], f"{CAR}/数据/表现.csv"))
        self.assertEqual(body["row"]["metrics"], {"views": 120, "reach": None, "likes": 10,
                                                  "comments": None, "saves": None, "shares": None,
                                                  "profile_visits": None, "follows": None, "ctr": None,
                                                  "three_sec_rate": None, "completion_rate": None,
                                                  "avg_watch_sec": None})
        last = (self.paths["小车"] / "数据/表现.csv").read_text(encoding="utf-8").splitlines()[-1]
        self.assertEqual(last, "2026-09-26T10:00:00-07:00,rednote,示例创作者,120,,10,,,,,,,,,,,接口")
        metrics = self.get(f"/api/works/{self.car['id']}/metrics")[1]
        self.assertEqual(metrics["latest"]["metrics"]["views"], 120)
        status, body = self.post("/api/metrics/analyse", {"work": "小车先停好", "platform": "小红书",
                                                          "node": "D3", "conclusion": "标题太长",
                                                          "tags": "标题, 封面", "date": "2026-09-12"})
        self.assertEqual(status, 200, body)
        self.assertEqual(body["path"], f"{CAR}/数据/分析/2026-09-12-D3-2.md")   # never overwrites
        self.assertEqual(body["tags"], ["标题", "封面"])
        status, body = self.post("/api/metrics/retention", {"work": "小车先停好", "platform": "rednote",
                                                            "points": "0:100,3:70"})
        self.assertEqual(status, 200, body)
        self.assertEqual(body["points"], [[0, 100], [3, 70]])
        status, body = self.post("/api/metrics/account", {"platform": "rednote", "account": "示例创作者",
                                                          "followers": 230, "source": "主页"})
        self.assertEqual(status, 200, body)
        self.assertEqual(self.get("/api/metrics")[1]["accounts"][0]["latest"]["followers"], 230)
        tags = {item["tag"] for item in self.get("/api/metrics")[1]["tags"]}
        self.assertIn("标题", tags)

    def test_write_errors(self) -> None:
        self.assertEqual(self.post("/api/metrics/add", {"work": "x"}, "text/plain")[0], 415)
        status, body = self.post("/api/metrics/add", {"work": "小车先停好", "platform": "火星", "views": 1})
        self.assertEqual(status, 400)
        self.assertIn("平台", body["error"])
        status, body = self.post("/api/metrics/add", {"work": "小车先停好", "platform": "rednote"})
        self.assertEqual(status, 400)
        status, body = self.post("/api/metrics/add", {"work": "没有的", "platform": "rednote", "views": 1})
        self.assertEqual(status, 404)
        status, body = self.post("/api/metrics/add", {"work": "AK", "platform": "rednote", "views": 1})
        self.assertEqual(status, 404)
        self.assertIn("匹配到 4 件作品", body["error"])
        self.assertEqual(self.post("/api/metrics/analyse", {"work": "小车先停好"})[0], 400)
        self.assertEqual(self.post("/api/metrics/account", {"platform": "rednote", "followers": 1})[0], 400)
        self.assertEqual(self.post("/api/metrics/nope", {})[0], 409)


if __name__ == "__main__":
    unittest.main()
